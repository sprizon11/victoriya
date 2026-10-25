"""
Salary module checks: day/night shifts, rates that stay as paid, advances, exports, payslips, access.

Needs the API running on :8000 and MongoDB reachable. Uses weeks in 2020 and removes them afterwards.

    backend\\.venv\\Scripts\\python.exe backend\\test_salary.py
"""
import io
import os
import sys
import uuid
from datetime import date, timedelta
from pathlib import Path

import requests
from dotenv import load_dotenv
from openpyxl import load_workbook
from pymongo import MongoClient

ROOT = Path(__file__).parent
load_dotenv(ROOT / ".env")
BASE = os.environ.get("TEST_BASE_URL", "http://127.0.0.1:8000")
EMAIL = os.environ.get("ADMIN_EMAIL", "admin@example.com").strip('"')
PASSWORD = os.environ.get("ADMIN_PASSWORD", "admin123").strip('"')
mongo = MongoClient(os.environ["MONGO_URL"].strip('"'))[os.environ["DB_NAME"].strip('"')]

passed, failed = [], []


def check(name, cond, detail=""):
    (passed if cond else failed).append(name if cond else (name, detail))
    print(f"  {'PASS' if cond else 'FAIL'}  {name}   {'' if cond else str(detail)[:300]}")


def call(s, method, path, **kw):
    return s.request(method, f"{BASE}/api{path}", timeout=60, **kw)


tag = uuid.uuid4().hex[:6]
admin = requests.Session()
if call(admin, "post", "/auth/login", json={"email": EMAIL, "password": PASSWORD}).status_code != 200:
    print("cannot log in; is the API running?")
    sys.exit(1)
W1, W2 = "2020-01-04", "2020-01-11"        # both Saturdays
S = lambda path: f"/salary{path}"

print("--- departments and workers ---")
dept = call(admin, "post", S("/departments"), json={"name": f"Power table {tag}"}).json()
check("a department can be added", "id" in dept)
check("a duplicate department is refused", call(admin, "post", S("/departments"), json={"name": f"POWER TABLE {tag}"}).status_code == 400)
r = call(admin, "post", S("/workers"), json={"name": f"Ravi {tag}", "designation": "Tailor", "emp_id": f"E{tag}", "department": dept["name"], "day_rate": 800})
check("a worker can be added", r.status_code == 200, r.text)
ravi = r.json()
check("a blank night rate follows the day rate", ravi["night_rate"] == 800 and ravi["night_same_as_day"] is True, ravi)
r = call(admin, "post", S("/workers"), json={"name": f"Kumar {tag}", "designation": "Helper", "department": dept["name"], "day_rate": 500, "night_rate": 600})
kumar = r.json()
check("a different night rate is kept", kumar["day_rate"] == 500 and kumar["night_rate"] == 600)
check("a duplicate employee id is refused", call(admin, "post", S("/workers"), json={"name": "X", "emp_id": f"e{tag}", "day_rate": 1}).status_code == 400)
check("an unknown department is refused", call(admin, "post", S("/workers"), json={"name": "Y", "department": "Nowhere", "day_rate": 1}).status_code == 400)
check("a negative rate is refused", call(admin, "post", S("/workers"), json={"name": "Z", "day_rate": -5}).status_code == 422)
check("a department with workers cannot be deleted", call(admin, "delete", S(f"/departments/{dept['id']}")).status_code == 400)
check("the worker list filters by department", len(call(admin, "get", S("/workers"), params={"department": dept["name"]}).json()) == 2)
check("...and searches by name", [w["name"] for w in call(admin, "get", S("/workers"), params={"q": f"kumar {tag}"}).json()] == [f"Kumar {tag}"])
check("a worker can be edited", call(admin, "put", S(f"/workers/{kumar['id']}"), json={"designation": "Senior helper"}).json()["designation"] == "Senior helper")

print("\n--- the pay week ---")
w = call(admin, "get", S("/week"), params={"date": "2020-01-08"}).json()
check("any day maps to the Saturday that starts its week", w["week_start"] == W1 and w["week_end"] == "2020-01-10", w["week_start"])
check("a Saturday maps to itself and a Friday to the week before", call(admin, "get", S("/week"), params={"date": W1}).json()["week_start"] == W1
      and call(admin, "get", S("/week"), params={"date": "2020-01-10"}).json()["week_start"] == W1)
check("the week lists seven dates, Saturday first", w["days"][0] == W1 and w["days"][6] == "2020-01-10" and w["day_names"][0] == "Sat")
check("a week that does not start on a Saturday is refused", call(admin, "get", S("/week"), params={"start": "2020-01-06"}).status_code == 400)
check("a bad date is refused", call(admin, "get", S("/week"), params={"date": "banana"}).status_code == 400)

print("\n--- the calculator maths ---")
body = {"day_shifts": [1, 1, 1, 1, 1, 1, 0], "night_shifts": [0, 0, 0, 0, 0, 0, 1.5], "advance": 500}
r = call(admin, "put", S(f"/week/{W1}/worker/{ravi['id']}"), json=body)
check("a week is saved", r.status_code == 200, r.text)
e = r.json()
check("6 day shifts x 800 = 4,800", e["day_total"] == 6 and e["day_wage"] == 4800.0, e)
check("1.5 night shifts x 800 = 1,200 (night rate defaults to the day rate)", e["night_total"] == 1.5 and e["night_wage"] == 1200.0)
check("gross 6,000, advance 500, net 5,500, and no ESI/PF anywhere", e["gross"] == 6000.0 and e["advance"] == 500.0 and e["net"] == 5500.0 and not any(k.lower() in ("esi", "pf", "esi_pf", "esi_pf_deduction") or k.lower().startswith(("esi_", "pf_")) for k in e))
r = call(admin, "put", S(f"/week/{W1}/worker/{kumar['id']}"), json={"day_shifts": [0, 0, 0, 0, 0, 0, 0], "night_shifts": [1, 1, 1, 1, 1, 0.5, 0], "advance": 0})
check("night shifts use the night rate: 5.5 x 600 = 3,300", r.json()["night_wage"] == 3300.0 and r.json()["net"] == 3300.0, r.text)
r = call(admin, "put", S(f"/week/{W2}/worker/{ravi['id']}"), json={"day_shifts": [0.5, 0, 0, 0, 0, 0, 0], "night_shifts": [0] * 7, "advance": 0})
cheap = call(admin, "put", S(f"/worker/{ravi['id']}"), json={}) if False else None
check("half a shift is allowed (0.5 x 800 = 400)", r.status_code == 200 and r.json()["day_wage"] == 400.0)
for label, b in [("0.3 of a shift", {"day_shifts": [0.3, 0, 0, 0, 0, 0, 0], "night_shifts": [0] * 7}),
                 ("3.5 shifts in a day", {"day_shifts": [3.5, 0, 0, 0, 0, 0, 0], "night_shifts": [0] * 7}),
                 ("a negative shift", {"day_shifts": [-1, 0, 0, 0, 0, 0, 0], "night_shifts": [0] * 7}),
                 ("day and night over 3 together", {"day_shifts": [2, 0, 0, 0, 0, 0, 0], "night_shifts": [2, 0, 0, 0, 0, 0, 0]})]:
    check(f"{label} is refused", call(admin, "put", S(f"/week/{W2}/worker/{kumar['id']}"), json=b).status_code == 400, label)
check("six days instead of seven is refused", call(admin, "put", S(f"/week/{W2}/worker/{kumar['id']}"), json={"day_shifts": [1] * 6, "night_shifts": [0] * 7}).status_code == 422)
r = call(admin, "put", S(f"/week/{W2}/worker/{kumar['id']}"), json={"day_shifts": [1, 0, 0, 0, 0, 0, 0], "night_shifts": [0] * 7, "advance": 9999})
check("an advance bigger than the wage is refused with the amounts", r.status_code == 400 and "more than the wage" in r.text, r.text)
r = call(admin, "put", S(f"/week/{W2}/worker/{kumar['id']}"), json={"day_shifts": [0.5, 0, 0, 0, 0, 0, 0], "night_shifts": [0] * 7, "advance": 0})
mongo.salary_workers.update_one({"_id": __import__("bson").ObjectId(kumar["id"])}, {"$set": {"day_rate": 333.33, "night_rate": 333.33}})
r = call(admin, "put", S(f"/week/{W2}/worker/{kumar['id']}"), json={"day_shifts": [0.5, 0, 0, 0, 0, 0, 0], "night_shifts": [0] * 7, "advance": 0, "reprice": True})
check("pennies round half up: 0.5 x 333.33 = 166.67 (not 166.66)", r.json()["day_wage"] == 166.67, r.json()["day_wage"])
call(admin, "put", S(f"/week/{W2}/worker/{kumar['id']}"), json={"day_shifts": [0] * 7, "night_shifts": [0] * 7, "advance": 0, "reprice": True})
mongo.salary_workers.update_one({"_id": __import__("bson").ObjectId(kumar["id"])}, {"$set": {"day_rate": 500, "night_rate": 600}})

print("\n--- rates stay as paid ---")
call(admin, "put", S(f"/workers/{ravi['id']}"), json={"day_rate": 1000})
r = call(admin, "put", S(f"/week/{W1}/worker/{ravi['id']}"), json={**body, "advance": 100})
check("editing a saved week keeps the rate it was paid at", r.json()["day_rate"] == 800 and r.json()["day_wage"] == 4800.0 and r.json()["net"] == 5900.0, r.json())
r = call(admin, "put", S(f"/week/{W1}/worker/{ravi['id']}"), json={**body, "advance": 100, "reprice": True})
check("...unless the user asks to use today's rate", r.json()["day_rate"] == 1000 and r.json()["day_wage"] == 6000.0, r.json())
r = call(admin, "put", S(f"/week/{W1}/worker/{ravi['id']}"), json={**body, "advance": 100})
check("...and after that the new rate is what stays", r.json()["day_rate"] == 1000)
check("saving twice never duplicates the row", mongo.salary_entries.count_documents({"week_start": W1, "worker_id": ravi["id"]}) == 1)
check("each re-save leaves a history trail", len(mongo.salary_entries.find_one({"week_start": W1, "worker_id": ravi["id"]})["history"]) >= 3)

print("\n--- the week page, history ---")
wk = call(admin, "get", S("/week"), params={"start": W1}).json()
check("both workers are on the week", [x["worker_name"] for x in wk["entries"]] == [f"Kumar {tag}", f"Ravi {tag}"])
check("totals add up", wk["totals"]["workers"] == 2 and wk["totals"]["gross"] == round(7500 + 3300, 2) and wk["totals"]["net"] == round(7500 - 100 + 3300, 2), wk["totals"])
h = call(admin, "get", S("/weeks")).json()
mine = [x for x in h["items"] if x["week_start"] in (W1, W2)]
check("history lists the weeks newest first with totals", [x["week_start"] for x in mine] == [W2, W1] and mine[1]["workers"] == 2, mine)

print("\n--- Excel and payslips ---")
x = call(admin, "get", S("/export"), params={"week": W1})
check("the week exports to Excel", x.status_code == 200 and x.content[:2] == b"PK")
ws = load_workbook(io.BytesIO(x.content)).active
head = [c.value for c in ws[1]]
check("the sheet has day and night shift columns for each weekday", [h for h in head if str(h).startswith("Day ")][:2] == ["Day Sat", "Day Sun"] and "Night Fri" in head and "Night shifts" in head, head)
check("...and no ESI/PF or piece-rate columns", not any(str(h).lower() in ("esi", "pf", "esi/pf", "piece rate", "pieces") or "piece" in str(h).lower() for h in head))
rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
ravi_row = next(r for r in rows if r[3] == f"Ravi {tag}")
check("a worker's row carries the figures", ravi_row[head.index("Gross")] == 7500.0 and ravi_row[head.index("Net pay")] == 7400.0, ravi_row)
total_row = next(r for r in rows if r[3] == "TOTAL")
check("a totals row sums the week", total_row[head.index("Gross")] == 10800.0 and total_row[head.index("Net pay")] == 10700.0, total_row)
xa = call(admin, "get", S("/export"), params={"all": True})
wb = load_workbook(io.BytesIO(xa.content))
check("'all weeks' has a sheet per purpose", xa.status_code == 200 and "Weeks" in wb.sheetnames)
check("an empty week cannot be exported", call(admin, "get", S("/export"), params={"week": "2019-12-28"}).status_code == 404)
check("the export needs a week or all", call(admin, "get", S("/export")).status_code == 400)
p = call(admin, "get", S(f"/payslip/{W1}/{ravi['id']}"))
check("one payslip as a PDF", p.status_code == 200 and p.content[:4] == b"%PDF")
pa = call(admin, "get", S(f"/payslips/{W1}"))
check("every payslip of the week in one PDF", pa.status_code == 200 and pa.content[:4] == b"%PDF" and len(pa.content) > len(p.content))
check("a payslip for an unsaved week is 404", call(admin, "get", S(f"/payslip/{W2}/000000000000000000000000")).status_code == 404)

print("\n--- access and safety ---")
clerk_email, acct_email, acct2_email = f"s-{tag}@example.com", f"sa-{tag}@example.com", f"sb-{tag}@example.com"
call(admin, "post", "/users", json={"email": clerk_email, "password": "Sal!" + tag, "name": "S Clerk", "role": "data_entry"})
call(admin, "post", "/users", json={"email": acct_email, "password": "Sal!" + tag, "name": "S Acct", "role": "accounts", "can_view_financial": True})
call(admin, "post", "/users", json={"email": acct2_email, "password": "Sal!" + tag, "name": "S Acct2", "role": "accounts", "can_view_financial": False})
sess = {}
for k, em in (("clerk", clerk_email), ("acct", acct_email), ("acct2", acct2_email)):
    sess[k] = requests.Session()
    call(sess[k], "post", "/auth/login", json={"email": em, "password": "Sal!" + tag})
check("data entry cannot see salary at all", call(sess["clerk"], "get", S("/week")).status_code == 403 and call(sess["clerk"], "get", S("/workers")).status_code == 403)
check("accounts with financial access can use it", call(sess["acct"], "get", S("/week")).status_code == 200)
check("accounts without financial access cannot", call(sess["acct2"], "get", S("/week")).status_code == 403)
check("an anonymous request is refused", requests.get(f"{BASE}/api/salary/week", timeout=30).status_code == 401)
check("accounts cannot clear a week", call(sess["acct"], "post", S(f"/week/{W1}/clear"), json={"confirm": W1}).status_code == 403)
check("a worker with saved weeks cannot be deleted", call(admin, "delete", S(f"/workers/{ravi['id']}")).status_code == 400)
call(admin, "put", S(f"/workers/{kumar['id']}"), json={"active": False})
r = call(admin, "put", S(f"/week/2020-01-18/worker/{kumar['id']}"), json={"day_shifts": [1] * 7, "night_shifts": [0] * 7})
check("a switched-off worker cannot be newly paid", r.status_code == 400 and "switched off" in r.text, r.text)
call(admin, "put", S(f"/workers/{kumar['id']}"), json={"active": True})
check("pay changes are in the activity log", call(admin, "get", "/audit-logs", params={"entity_type": "salary_entry", "q": f"Ravi {tag}"}).json()["total"] >= 3)
check("removing one worker from a week works", call(admin, "delete", S(f"/week/{W2}/worker/{ravi['id']}")).status_code == 200)
check("clearing a week needs the date typed back", call(admin, "post", S(f"/week/{W1}/clear"), json={"confirm": "yes"}).status_code == 400)
cl = call(admin, "post", S(f"/week/{W1}/clear"), json={"confirm": W1})
check("an administrator can clear a week", cl.status_code == 200 and cl.json()["cleared"] == 2, cl.text)
check("...a copy is kept", mongo.salary_cleared.count_documents({"week_start": W1}) == 1)
check("...and the week is empty", call(admin, "get", S("/week"), params={"start": W1}).json()["entries"] == [])

print("\n--- time-based pay (salary per day, hours per day) ---")
W3 = "2020-01-25"
check("a time-based worker needs hours per day", call(admin, "post", S("/workers"), json={"name": f"NoHours {tag}", "day_rate": 800, "pay_type": "time"}).status_code == 400)
check("an unknown pay type is refused", call(admin, "post", S("/workers"), json={"name": f"Odd {tag}", "day_rate": 800, "pay_type": "weekly"}).status_code == 422)
check("zero hours per day is refused", call(admin, "post", S("/workers"), json={"name": f"Zero {tag}", "day_rate": 800, "pay_type": "time", "hours_per_day": 0}).status_code == 422)
r = call(admin, "post", S("/workers"), json={"name": f"Timed {tag}", "department": dept["name"], "day_rate": 800, "night_rate": 960, "pay_type": "time", "hours_per_day": 8})
timed = r.json()
check("a time-based worker is created with hours per day", r.status_code == 200 and timed["pay_type"] == "time" and timed["hours_per_day"] == 8, r.text)
check("older and shift workers read as pay_type shift", ravi["pay_type"] == "shift" and ravi["hours_per_day"] is None)
r = call(admin, "put", S(f"/week/{W3}/worker/{timed['id']}"), json={"day_shifts": [8, 8, 8, 8, 8, 8, 4], "night_shifts": [0, 0, 0, 0, 0, 0, 2.5], "advance": 100})
t = r.json()
check("hours / 8 x day salary: 52 h x 800 / 8 = 5,200", r.status_code == 200 and t["day_wage"] == 5200.0, r.text)
check("night hours use the night salary: 2.5 h x 960 / 8 = 300", t["night_wage"] == 300.0 and t["gross"] == 5500.0 and t["net"] == 5400.0, t)
check("the entry records it as time-based with its hours per day", t["pay_type"] == "time" and t["hours_per_day"] == 8 and t["day_total"] == 52)
check("totals count hours as days (52 h / 8 = 6.5)", t["day_days"] == 6.5 and call(admin, "get", S("/week"), params={"start": W3}).json()["totals"]["day_shifts"] == 6.5)
odd = call(admin, "post", S("/workers"), json={"name": f"Nine {tag}", "day_rate": 500, "pay_type": "time", "hours_per_day": 9}).json()
r = call(admin, "put", S(f"/week/{W3}/worker/{odd['id']}"), json={"day_shifts": [7, 0, 0, 0, 0, 0, 0], "night_shifts": [0] * 7})
check("rounded once at the end: 7 h x 500 / 9 = 388.89 (not 7 x 55.56)", r.json()["day_wage"] == 388.89, r.text)
for label, b in [("0.3 of an hour", {"day_shifts": [0.3, 0, 0, 0, 0, 0, 0], "night_shifts": [0] * 7}),
                 ("25 hours in a day", {"day_shifts": [25, 0, 0, 0, 0, 0, 0], "night_shifts": [0] * 7}),
                 ("day and night over 24 hours together", {"day_shifts": [14, 0, 0, 0, 0, 0, 0], "night_shifts": [11, 0, 0, 0, 0, 0, 0]})]:
    check(f"{label} is refused", call(admin, "put", S(f"/week/{W3}/worker/{timed['id']}"), json=b).status_code == 400)
check("quarter hours are allowed (8.25 h)", call(admin, "put", S(f"/week/{W3}/worker/{odd['id']}"), json={"day_shifts": [8.25, 0, 0, 0, 0, 0, 0], "night_shifts": [0] * 7}).status_code == 200)
call(admin, "put", S(f"/workers/{timed['id']}"), json={"day_rate": 1000})
r = call(admin, "put", S(f"/week/{W3}/worker/{timed['id']}"), json={"day_shifts": [8, 8, 8, 8, 8, 8, 4], "night_shifts": [0, 0, 0, 0, 0, 0, 2.5], "advance": 100})
check("a saved week keeps the salary it was paid at", r.json()["day_wage"] == 5200.0 and r.json()["day_rate"] == 800)
r = call(admin, "put", S(f"/week/{W3}/worker/{timed['id']}"), json={"day_shifts": [8, 8, 8, 8, 8, 8, 4], "night_shifts": [0, 0, 0, 0, 0, 0, 2.5], "advance": 100, "reprice": True})
check("...unless today's salary is asked for (52 h x 1000 / 8 = 6,500)", r.json()["day_wage"] == 6500.0)
check("switching a worker to shift pay is allowed", call(admin, "put", S(f"/workers/{odd['id']}"), json={"pay_type": "shift"}).json()["pay_type"] == "shift")
check("...and back to time needs the hours already stored", call(admin, "put", S(f"/workers/{odd['id']}"), json={"pay_type": "time"}).json()["hours_per_day"] == 9)
x = load_workbook(io.BytesIO(call(admin, "get", S("/export"), params={"week": W3}).content)).active
head = [c.value for c in x[1]]
check("the Excel sheet says whether pay is by hours or by shift", "Pay type" in head and "Hours per day" in head and x.cell(row=2, column=head.index("Pay type") + 1).value in ("By hours", "By shift"), head)
check("the payslip PDF works for a time-based worker", call(admin, "get", S(f"/payslip/{W3}/{timed['id']}")).content[:4] == b"%PDF")

print("\n--- cleanup ---")
mongo.salary_entries.delete_many({"week_start": {"$in": [W1, W2, W3, "2020-01-18"]}})
mongo.salary_cleared.delete_many({"week_start": W1})
mongo.salary_workers.delete_many({"name": {"$regex": f"{tag}$"}})
mongo.salary_departments.delete_many({"name": {"$regex": f"{tag}$"}})
for u in call(admin, "get", "/users").json():
    if u["email"] in (clerk_email, acct_email, acct2_email):
        call(admin, "delete", f"/users/{u['id']}")
print("  removed fixtures")

print("\n" + "=" * 60)
print(f"  {len(passed)} passed, {len(failed)} failed")
print("=" * 60)
sys.exit(1 if failed else 0)
