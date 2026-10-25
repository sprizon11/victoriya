"""Expense bills: service, maintenance, accessories and anything else the unit pays for.

A bill is never deleted: it is voided with a reason. Bills attributed to a lot feed
that lot's costing; the rest are overhead. Bill photos/PDFs are kept in MongoDB
GridFS because the hosting disk is ephemeral.
"""
import io
import re
from typing import Optional

from fastapi import APIRouter, HTTPException, Query, Request, UploadFile, File
from fastapi.responses import StreamingResponse, Response
from motor.motor_asyncio import AsyncIOMotorGridFSBucket
from pydantic import BaseModel, Field

from core import (db, ObjectId, datetime, timezone, get_current_user, check_role, serialize_doc, get_fy_suffix)
from audit import note
import approvals
from approvals import record_auto, ApprovalExpired

router = APIRouter(prefix="/api")
MAX_BYTES = 5 * 1024 * 1024
PAYMENT_MODES = ["cash", "upi", "bank", "cheque", "other"]


def _now():
    return datetime.now(timezone.utc).isoformat()


def sniff(data: bytes):
    """Type from the file's own bytes: a renamed .exe must not pass as a bill."""
    if data[:4] == b"%PDF":
        return "application/pdf"
    if data[:3] == b"\xff\xd8\xff":
        return "image/jpeg"
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "image/png"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    return None


async def next_voucher() -> str:
    fy = get_fy_suffix()
    res = await db.counters.find_one_and_update({"name": f"exp_{fy}"}, {"$inc": {"value": 1}}, upsert=True, return_document=True)
    return f"EXP{res['value']}/{fy}"


async def threshold() -> float:
    s = await db.system_settings.find_one({"key": "expense_approval_threshold"})
    try:
        return float(s["value"]) if s else 0.0
    except (TypeError, ValueError, KeyError):
        return 0.0


# ------------------------------------------------------------------ categories
class CategoryIn(BaseModel):
    name: str = Field(min_length=1, max_length=60)


@router.get("/expense-categories")
async def list_categories(request: Request):
    user = await get_current_user(request)
    check_role(user, ["admin", "accounts"])
    rows = await db.expense_categories.find().sort("name", 1).to_list(200)
    return [serialize_doc(r) for r in rows]


@router.post("/expense-categories")
async def create_category(data: CategoryIn, request: Request):
    user = await get_current_user(request)
    check_role(user, ["admin"])
    name = data.name.strip()
    if await db.expense_categories.find_one({"name": {"$regex": f"^{re.escape(name)}$", "$options": "i"}}):
        raise HTTPException(400, f"A category called '{name}' already exists")
    doc = {"name": name, "active": True, "created_at": _now()}
    res = await db.expense_categories.insert_one(doc)
    note(request, entity_type="expense_category", entity_id=str(res.inserted_id), label=name, summary=f"Created expense category {name}")
    return {"id": str(res.inserted_id), "name": name, "active": True}


@router.put("/expense-categories/{cat_id}")
async def update_category(cat_id: str, request: Request):
    user = await get_current_user(request)
    check_role(user, ["admin"])
    body = await request.json()
    if not ObjectId.is_valid(cat_id):
        raise HTTPException(404, "Category not found")
    cat = await db.expense_categories.find_one({"_id": ObjectId(cat_id)})
    if not cat:
        raise HTTPException(404, "Category not found")
    update = {}
    if "name" in body and str(body["name"]).strip():
        update["name"] = str(body["name"]).strip()[:60]
    if "active" in body:
        update["active"] = bool(body["active"])
    if not update:
        raise HTTPException(400, "No fields to update")
    await db.expense_categories.update_one({"_id": cat["_id"]}, {"$set": update})
    if "name" in update:
        await db.expenses.update_many({"category_id": cat_id}, {"$set": {"category_name": update["name"]}})
    note(request, entity_type="expense_category", entity_id=cat_id, label=cat["name"], summary=f"Updated expense category {cat['name']}",
         before={k: cat.get(k) for k in update}, after=update)
    return serialize_doc(await db.expense_categories.find_one({"_id": cat["_id"]}))


@router.delete("/expense-categories/{cat_id}")
async def delete_category(cat_id: str, request: Request):
    user = await get_current_user(request)
    check_role(user, ["admin"])
    if not ObjectId.is_valid(cat_id):
        raise HTTPException(404, "Category not found")
    if await db.expenses.count_documents({"category_id": cat_id}):
        raise HTTPException(400, "Bills exist in this category. Switch it off instead of deleting it.")
    cat = await db.expense_categories.find_one({"_id": ObjectId(cat_id)})
    if not cat:
        raise HTTPException(404, "Category not found")
    note(request, entity_type="expense_category", entity_id=cat_id, label=cat["name"], summary=f"Deleted expense category {cat['name']}")
    await db.expense_categories.delete_one({"_id": cat["_id"]})
    return {"message": "Deleted"}


# ------------------------------------------------------------------ expenses
class ExpenseIn(BaseModel):
    date: str = Field(min_length=10, max_length=10)
    category_id: str
    vendor: str = Field(min_length=1, max_length=120)
    bill_no: str = Field(default="", max_length=60)
    bill_date: Optional[str] = None
    description: str = Field(default="", max_length=500)
    amount: float = Field(gt=0, le=100_000_000)
    gst_amount: float = Field(default=0, ge=0, le=100_000_000)
    payment_status: str = "unpaid"
    paid_amount: float = Field(default=0, ge=0)
    payment_mode: str = ""
    paid_on: Optional[str] = None
    lot_id: Optional[str] = None
    machine_ref: str = Field(default="", max_length=80)
    allow_duplicate: bool = False
    duplicate_reason: str = Field(default="", max_length=200)


class ExpenseUpdate(BaseModel):
    category_id: Optional[str] = None
    vendor: Optional[str] = Field(default=None, min_length=1, max_length=120)
    bill_no: Optional[str] = Field(default=None, max_length=60)
    bill_date: Optional[str] = None
    description: Optional[str] = Field(default=None, max_length=500)
    amount: Optional[float] = Field(default=None, gt=0, le=100_000_000)
    gst_amount: Optional[float] = Field(default=None, ge=0, le=100_000_000)
    lot_id: Optional[str] = None
    machine_ref: Optional[str] = Field(default=None, max_length=80)


class PayIn(BaseModel):
    amount: float = Field(gt=0, le=100_000_000)
    mode: str = "cash"
    date: Optional[str] = None


async def _expense(expense_id):
    if not ObjectId.is_valid(expense_id):
        raise HTTPException(404, "Bill not found")
    e = await db.expenses.find_one({"_id": ObjectId(expense_id)})
    if not e:
        raise HTTPException(404, "Bill not found")
    return e


def derive_payment(e):
    paid = round(e.get("paid_amount", 0), 2)
    total = e["total"]
    e["payment_status"] = "paid" if paid >= total - 0.005 else ("partial" if paid > 0 else "unpaid")
    e["balance"] = round(max(0, total - paid), 2)
    return e


def present(e):
    d = serialize_doc(dict(e))
    d["balance"] = round(max(0, d["total"] - d.get("paid_amount", 0)), 2)
    d["has_attachment"] = bool(d.get("attachment"))
    return d


async def _lot_label(lot_id):
    if not lot_id:
        return ""
    lot = await db.lots.find_one({"_id": ObjectId(lot_id)}) if ObjectId.is_valid(lot_id) else None
    if not lot:
        raise HTTPException(400, "Lot not found")
    return lot["lot_number"]


@router.post("/expenses")
async def create_expense(data: ExpenseIn, request: Request):
    user = await get_current_user(request)
    check_role(user, ["admin", "accounts"])
    if not ObjectId.is_valid(data.category_id):
        raise HTTPException(400, "Choose a category")
    cat = await db.expense_categories.find_one({"_id": ObjectId(data.category_id)})
    if not cat or not cat.get("active", True):
        raise HTTPException(400, "Category not found or switched off")
    if data.payment_mode and data.payment_mode not in PAYMENT_MODES:
        raise HTTPException(400, f"Payment mode must be one of {', '.join(PAYMENT_MODES)}")
    vendor = data.vendor.strip()
    if data.bill_no.strip():
        dup = await db.expenses.find_one({"vendor": {"$regex": f"^{re.escape(vendor)}$", "$options": "i"},
                                          "bill_no": {"$regex": f"^{re.escape(data.bill_no.strip())}$", "$options": "i"},
                                          "status": {"$ne": "void"}})
        if dup and not (data.allow_duplicate and data.duplicate_reason.strip()):
            raise HTTPException(400, f"{vendor} bill {data.bill_no} is already recorded as {dup['voucher_no']}. "
                                     "If it really is a second bill, tick 'allow duplicate' and say why.")
    lot_number = await _lot_label(data.lot_id)
    total = round(data.amount + data.gst_amount, 2)
    paid = round(min(data.paid_amount, total), 2)
    limit = await threshold()
    needs_ok = limit > 0 and total > limit and (await approvals.policies()).get("expense_over_limit")
    status = "approved"
    if needs_ok and user["role"] != "admin":
        status = "pending_approval"
    doc = {"voucher_no": await next_voucher(), "date": data.date, "category_id": data.category_id, "category_name": cat["name"],
           "vendor": vendor, "bill_no": data.bill_no.strip(), "bill_date": data.bill_date or data.date,
           "description": data.description.strip(), "amount": round(data.amount, 2), "gst_amount": round(data.gst_amount, 2),
           "total": total, "paid_amount": paid, "payment_mode": data.payment_mode if paid else "",
           "paid_on": (data.paid_on or data.date) if paid else "", "lot_id": data.lot_id or "", "lot_number": lot_number,
           "machine_ref": data.machine_ref.strip(), "status": status, "duplicate_reason": data.duplicate_reason.strip() if data.allow_duplicate else "",
           "created_by": user["id"], "created_by_name": user.get("name", ""), "created_at": _now()}
    derive_payment(doc)
    res = await db.expenses.insert_one(doc)
    doc["_id"] = res.inserted_id
    eid = str(res.inserted_id)
    approval_id = None
    if status == "pending_approval":
        from approvals import create_request
        req = await create_request(request, user, "expense_over_limit", eid, f"Bill above the approval limit ({limit:g})", {})
        approval_id = req["id"]
    elif needs_ok:
        approval_id = await record_auto(request, user, "expense_over_limit", "expense", eid, doc["voucher_no"],
                                        f"{doc['voucher_no']} {vendor} {total:g}")
    note(request, entity_type="expense", entity_id=eid, label=f"{doc['voucher_no']} {vendor}", approval_id=approval_id,
         summary=f"Recorded {cat['name']} bill {total:g} from {vendor}" + (" (waiting for approval)" if status == "pending_approval" else ""),
         after={"total": total, "category": cat["name"], "lot": lot_number})
    return present(doc)


@router.get("/expenses")
async def list_expenses(request: Request, date_from: Optional[str] = None, date_to: Optional[str] = None,
                        category_id: Optional[str] = None, vendor: Optional[str] = None, payment_status: Optional[str] = None,
                        lot_id: Optional[str] = None, status: Optional[str] = None, q: Optional[str] = None,
                        page: int = Query(1, ge=1), page_size: int = Query(50, ge=1, le=200)):
    user = await get_current_user(request)
    check_role(user, ["admin", "accounts"])
    query = _filter(date_from, date_to, category_id, vendor, payment_status, lot_id, status, q)
    total = await db.expenses.count_documents(query)
    rows = await db.expenses.find(query).sort([("date", -1), ("created_at", -1)]).skip((page - 1) * page_size).limit(page_size).to_list(page_size)
    sums = {"total": 0.0, "paid": 0.0, "unpaid": 0.0}
    async for s in db.expenses.aggregate([{"$match": {**query, "status": {"$nin": ["void", "rejected"]}}},
                                          {"$group": {"_id": None, "t": {"$sum": "$total"}, "p": {"$sum": "$paid_amount"}}}]):
        sums = {"total": round(s["t"], 2), "paid": round(s["p"], 2), "unpaid": round(s["t"] - s["p"], 2)}
    return {"items": [present(r) for r in rows], "count": total, "page": page, "pages": max(1, -(-total // page_size)), "totals": sums}


def _filter(date_from=None, date_to=None, category_id=None, vendor=None, payment_status=None, lot_id=None, status=None, q=None):
    query = {}
    if date_from or date_to:
        query["date"] = {**({"$gte": date_from} if date_from else {}), **({"$lte": date_to} if date_to else {})}
    if category_id:
        query["category_id"] = category_id
    if vendor:
        query["vendor"] = {"$regex": re.escape(vendor), "$options": "i"}
    if payment_status:
        query["payment_status"] = payment_status
    if lot_id:
        query["lot_id"] = lot_id
    if status:
        query["status"] = status
    if q:
        rx = {"$regex": re.escape(q), "$options": "i"}
        query["$or"] = [{"vendor": rx}, {"bill_no": rx}, {"description": rx}, {"voucher_no": rx}, {"machine_ref": rx}]
    return query


@router.get("/expenses/summary")
async def expense_summary(request: Request, month: Optional[str] = None):
    user = await get_current_user(request)
    check_role(user, ["admin", "accounts"])
    month = month or _now()[:7]
    live = {"status": {"$nin": ["void", "rejected", "pending_approval"]}}
    by_cat = []
    async for g in db.expenses.aggregate([{"$match": {**live, "date": {"$regex": f"^{re.escape(month)}"}}},
                                          {"$group": {"_id": "$category_name", "total": {"$sum": "$total"}, "n": {"$sum": 1}}},
                                          {"$sort": {"total": -1}}]):
        by_cat.append({"category": g["_id"], "total": round(g["total"], 2), "count": g["n"]})
    today = datetime.now(timezone.utc).date()
    aging = {"0-30": 0.0, "31-60": 0.0, "60+": 0.0}
    async for e in db.expenses.find({**live, "payment_status": {"$ne": "paid"}}):
        try:
            age = (today - datetime.strptime(e.get("bill_date") or e["date"], "%Y-%m-%d").date()).days
        except ValueError:
            age = 0
        bal = max(0, e["total"] - e.get("paid_amount", 0))
        aging["0-30" if age <= 30 else "31-60" if age <= 60 else "60+"] += bal
    return {"month": month, "by_category": by_cat, "month_total": round(sum(c["total"] for c in by_cat), 2),
            "unpaid_aging": {k: round(v, 2) for k, v in aging.items()}, "unpaid_total": round(sum(aging.values()), 2)}


@router.get("/expenses/export")
async def export_expenses(request: Request, date_from: Optional[str] = None, date_to: Optional[str] = None,
                          category_id: Optional[str] = None, vendor: Optional[str] = None, payment_status: Optional[str] = None,
                          lot_id: Optional[str] = None, status: Optional[str] = None, q: Optional[str] = None):
    user = await get_current_user(request)
    check_role(user, ["admin", "accounts"])
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"
    head = ["Voucher", "Date", "Category", "Vendor", "Bill no", "Description", "Lot", "Amount", "GST", "Total", "Paid", "Balance", "Status"]
    ws.append(head)
    for i in range(1, len(head) + 1):
        ws.cell(row=1, column=i).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=1, column=i).fill = PatternFill(start_color="294F96", end_color="294F96", fill_type="solid")
    tot = 0.0
    for e in await db.expenses.find(_filter(date_from, date_to, category_id, vendor, payment_status, lot_id, status, q)).sort("date", 1).to_list(20000):
        bal = round(max(0, e["total"] - e.get("paid_amount", 0)), 2)
        # a leading = + - @ would make Excel run the cell as a formula
        safe = lambda v: ("'" + v) if isinstance(v, str) and v[:1] in "=+-@" else v
        ws.append([e["voucher_no"], e["date"], e["category_name"], safe(e["vendor"]), safe(e.get("bill_no", "")), safe(e.get("description", "")),
                   e.get("lot_number", ""), e["amount"], e["gst_amount"], e["total"], e.get("paid_amount", 0), bal, e["status"]])
        if e["status"] not in ("void", "rejected"):
            tot += e["total"]
    ws.append([])
    ws.append(["", "", "", "", "", "", "TOTAL (excluding void/rejected)", "", "", round(tot, 2)])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=expenses.xlsx"})


@router.get("/expenses/{expense_id}")
async def get_expense(expense_id: str, request: Request):
    user = await get_current_user(request)
    check_role(user, ["admin", "accounts"])
    return present(await _expense(expense_id))


@router.put("/expenses/{expense_id}")
async def update_expense(expense_id: str, data: ExpenseUpdate, request: Request):
    user = await get_current_user(request)
    check_role(user, ["admin", "accounts"])
    e = await _expense(expense_id)
    if e["status"] in ("void", "rejected"):
        raise HTTPException(400, f"A {e['status']} bill cannot be edited")
    update = {k: v for k, v in data.model_dump().items() if v is not None}
    if not update:
        raise HTTPException(400, "No fields to update")
    if ("amount" in update or "gst_amount" in update) and user["role"] != "admin" and e["status"] == "approved":
        raise HTTPException(403, "Only an administrator can change the amount of an approved bill")
    if "category_id" in update:
        cat = await db.expense_categories.find_one({"_id": ObjectId(update["category_id"])}) if ObjectId.is_valid(update["category_id"]) else None
        if not cat:
            raise HTTPException(400, "Category not found")
        update["category_name"] = cat["name"]
    if "lot_id" in update:
        update["lot_number"] = await _lot_label(update["lot_id"])
    merged = {**e, **update}
    update["total"] = round(merged["amount"] + merged["gst_amount"], 2)
    merged["total"] = update["total"]
    update["payment_status"] = derive_payment({**merged})["payment_status"]
    await db.expenses.update_one({"_id": e["_id"]}, {"$set": update})
    note(request, entity_type="expense", entity_id=expense_id, label=f"{e['voucher_no']} {e['vendor']}",
         summary=f"Edited bill {e['voucher_no']}", before={k: e.get(k) for k in update}, after=update)
    return present(await _expense(expense_id))


@router.post("/expenses/{expense_id}/pay")
async def pay_expense(expense_id: str, data: PayIn, request: Request):
    user = await get_current_user(request)
    check_role(user, ["admin", "accounts"])
    e = await _expense(expense_id)
    if e["status"] != "approved":
        raise HTTPException(400, "Only an approved bill can be paid")
    if data.mode not in PAYMENT_MODES:
        raise HTTPException(400, f"Payment mode must be one of {', '.join(PAYMENT_MODES)}")
    balance = round(e["total"] - e.get("paid_amount", 0), 2)
    if data.amount > balance + 0.005:
        raise HTTPException(400, f"Only {balance:g} is still owed on this bill")
    paid = round(e.get("paid_amount", 0) + data.amount, 2)
    st = derive_payment({**e, "paid_amount": paid})["payment_status"]
    await db.expenses.update_one({"_id": e["_id"]}, {"$set": {"paid_amount": paid, "payment_status": st, "payment_mode": data.mode,
                                                              "paid_on": data.date or _now()[:10]},
                                                    "$push": {"payments": {"amount": data.amount, "mode": data.mode, "date": data.date or _now()[:10], "by": user.get("name", "")}}})
    note(request, entity_type="expense", entity_id=expense_id, label=f"{e['voucher_no']} {e['vendor']}",
         summary=f"Paid {data.amount:g} ({data.mode}) on {e['voucher_no']}", after={"paid_amount": paid, "status": st})
    return present(await _expense(expense_id))


@router.post("/expenses/{expense_id}/void")
async def void_expense(expense_id: str, request: Request):
    user = await get_current_user(request)
    check_role(user, ["admin"])
    body = await request.json()
    reason = str(body.get("reason", "")).strip()
    if not reason:
        raise HTTPException(400, "A reason is required to void a bill")
    e = await _expense(expense_id)
    if e["status"] == "void":
        raise HTTPException(400, "This bill is already void")
    await db.expenses.update_one({"_id": e["_id"]}, {"$set": {"status": "void", "void_reason": reason, "voided_by": user.get("name", ""), "voided_at": _now()}})
    note(request, entity_type="expense", entity_id=expense_id, label=f"{e['voucher_no']} {e['vendor']}",
         summary=f"Voided bill {e['voucher_no']}: {reason}", before={"status": e["status"], "total": e["total"]})
    return present(await _expense(expense_id))


@router.post("/expenses/{expense_id}/attachment")
async def upload_attachment(expense_id: str, request: Request, file: UploadFile = File(...)):
    user = await get_current_user(request)
    check_role(user, ["admin", "accounts"])
    e = await _expense(expense_id)
    data = await file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise HTTPException(400, "The file is larger than 5 MB")
    kind = sniff(data)
    if not kind:
        raise HTTPException(400, "Only PDF, JPEG, PNG or WebP files can be attached")
    bucket = AsyncIOMotorGridFSBucket(db, bucket_name="bills")
    if e.get("attachment"):
        try:
            await bucket.delete(ObjectId(e["attachment"]["file_id"]))
        except Exception:
            pass
    safe_name = re.sub(r"[^A-Za-z0-9._ -]", "_", file.filename or "bill")[:80]
    fid = await bucket.upload_from_stream(safe_name, data, metadata={"expense_id": expense_id, "content_type": kind})
    att = {"file_id": str(fid), "filename": safe_name, "content_type": kind, "size": len(data), "uploaded_at": _now()}
    await db.expenses.update_one({"_id": e["_id"]}, {"$set": {"attachment": att}})
    note(request, entity_type="expense", entity_id=expense_id, label=f"{e['voucher_no']} {e['vendor']}", summary=f"Attached {safe_name} to {e['voucher_no']}")
    return att


@router.get("/expenses/{expense_id}/attachment")
async def download_attachment(expense_id: str, request: Request):
    user = await get_current_user(request)
    check_role(user, ["admin", "accounts"])
    e = await _expense(expense_id)
    att = e.get("attachment")
    if not att:
        raise HTTPException(404, "No file attached")
    bucket = AsyncIOMotorGridFSBucket(db, bucket_name="bills")
    stream = await bucket.open_download_stream(ObjectId(att["file_id"]))
    data = await stream.read()
    return Response(content=data, media_type=att["content_type"],
                    headers={"Content-Disposition": f'inline; filename="{att["filename"]}"', "X-Content-Type-Options": "nosniff"})


# ------------------------------------------------------------------ approval: expense_over_limit
async def prepare_expense(user, expense_id, payload):
    e = await _expense(expense_id)
    if e["status"] != "pending_approval":
        raise HTTPException(400, "This bill is not waiting for approval")
    return f"{e['voucher_no']} · {e['vendor']} · {e['total']:g}", {}, {"total": e["total"]}


async def apply_expense(approval, reviewer, request):
    try:
        e = await _expense(approval["entity_id"])
    except HTTPException:
        raise ApprovalExpired("The bill no longer exists")
    if e["status"] != "pending_approval":
        raise ApprovalExpired(f"The bill is already {e['status']}")
    await db.expenses.update_one({"_id": e["_id"]}, {"$set": {"status": "approved", "approved_by": reviewer.get("name", ""), "approved_at": _now()}})
    return {"after": {"bill": e["voucher_no"], "status": "approved"}}


async def reject_expense(approval, reviewer, request):
    await db.expenses.update_one({"_id": ObjectId(approval["entity_id"]), "status": "pending_approval"},
                                 {"$set": {"status": "rejected", "rejected_by": reviewer.get("name", ""), "rejected_at": _now()}})


approvals.register("expense_over_limit", prepare=prepare_expense, apply=apply_expense, reject=reject_expense)
