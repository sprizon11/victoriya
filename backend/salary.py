"""Weekly salary for shift workers: a separate workforce from the piece-rate operators.

Modelled on the owner's reference calculator, with the changes they asked for: day and night
shifts are counted separately, there is no piece rate and no ESI/PF, workers can be added,
and everything exports to Excel.

A pay week runs Saturday to Friday. Saved weeks snapshot the worker's name and rates, so
raising a rate later never reprices a week that was already paid.
"""
import io
import re
from datetime import timedelta, date as date_cls
from decimal import Decimal, ROUND_HALF_UP
from typing import List, Literal, Optional

from fastapi import APIRouter, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from core import db, ObjectId, datetime, timezone, get_current_user, check_role, serialize_doc, COMPANY
from audit import note

router = APIRouter(prefix="/api/salary")
DAYS = ["Sat", "Sun", "Mon", "Tue", "Wed", "Thu", "Fri"]
HALF = Decimal("0.5")
QUARTER = Decimal("0.25")


def _now():
    return datetime.now(timezone.utc).isoformat()


async def salary_user(request: Request):
    user = await get_current_user(request)
    check_role(user, ["admin", "accounts"])
    if user["role"] != "admin" and not user.get("can_view_financial"):
        raise HTTPException(403, "Your account does not have permission to view financial data")
    return user


# ------------------------------------------------------------------ weeks
def parse_day(value: str) -> date_cls:
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        raise HTTPException(400, "Dates are written YYYY-MM-DD")


def week_start_of(d: date_cls) -> date_cls:
    """The Saturday on or before d (weekday(): Mon=0 ... Sat=5)."""
    return d - timedelta(days=(d.weekday() - 5) % 7)


def week_bounds(start: str):
    s = parse_day(start)
    if s.weekday() != 5:
        raise HTTPException(400, "Pay weeks start on a Saturday")
    return s.isoformat(), (s + timedelta(days=6)).isoformat()


# ------------------------------------------------------------------ maths
def money(x) -> float:
    return float(Decimal(str(x)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def clean_shifts(values, label, step=HALF, top=3):
    """Seven numbers, Sat..Fri. Shift workers: 1 = a full shift, 0.5 = half, 1.5 = a shift plus overtime.
    Time-based workers: the same seven boxes hold hours, in quarters, up to 24."""
    if not isinstance(values, list) or len(values) != 7:
        raise HTTPException(400, f"{label}: give a number for each of the 7 days")
    out = []
    for day, v in zip(DAYS, values):
        try:
            d = Decimal(str(v if v not in (None, "") else 0))
        except Exception:
            raise HTTPException(400, f"{label} on {day}: '{v}' is not a number")
        if d < 0 or d > top or (d % step) != 0:
            what = "0, 0.5, 1, 1.5, 2 … up to 3" if step == HALF else "hours in quarters (0, 0.25, 0.5 … up to 24)"
            raise HTTPException(400, f"{label} on {day}: use {what} (got {v})")
        out.append(float(d))
    return out


def compute(day_shifts, night_shifts, day_rate, night_rate, advance, hours_per_day=None):
    """Shift workers: units x rate. Time-based workers: hours / hours-per-day x the day's salary, so half a day's
    hours pay half the day's salary. Divided and rounded once at the end, never via a rounded hourly rate."""
    dt = sum(Decimal(str(x)) for x in day_shifts)
    nt = sum(Decimal(str(x)) for x in night_shifts)
    per = Decimal(str(hours_per_day)) if hours_per_day else Decimal(1)
    dw = (dt * Decimal(str(day_rate)) / per).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    nw = (nt * Decimal(str(night_rate)) / per).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    gross = dw + nw
    adv = Decimal(str(advance or 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return {"day_total": float(dt), "night_total": float(nt),
            "day_days": float((dt / per).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)),      # shifts, or hours expressed as days, so totals add up across both kinds
            "night_days": float((nt / per).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)),
            "day_wage": float(dw), "night_wage": float(nw),
            "gross": float(gross), "advance": float(adv), "net": float(gross - adv)}


# ------------------------------------------------------------------ departments + workers
class DeptIn(BaseModel):
    name: str = Field(min_length=1, max_length=60)


class WorkerIn(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    designation: str = Field(default="", max_length=80)
    emp_id: str = Field(default="", max_length=30)
    department: str = Field(default="", max_length=60)
    day_rate: float = Field(ge=0, le=100000)
    night_rate: Optional[float] = Field(default=None, ge=0, le=100000)      # blank = follows the day rate, even when that changes
    pay_type: Literal["shift", "time"] = "shift"          # time: the rates are the salary for a whole day of hours_per_day hours
    hours_per_day: Optional[float] = Field(default=None, gt=0, le=24)


class WorkerUpdate(BaseModel):
    name: Optional[str] = Field(default=None, min_length=1, max_length=100)
    designation: Optional[str] = Field(default=None, max_length=80)
    emp_id: Optional[str] = Field(default=None, max_length=30)
    department: Optional[str] = Field(default=None, max_length=60)
    day_rate: Optional[float] = Field(default=None, ge=0, le=100000)
    night_rate: Optional[float] = Field(default=None, ge=0, le=100000)
    pay_type: Optional[Literal["shift", "time"]] = None
    hours_per_day: Optional[float] = Field(default=None, gt=0, le=24)
    night_same_as_day: Optional[bool] = None      # True: forget a separate night rate and follow the day rate again
    active: Optional[bool] = None


@router.get("/departments")
async def list_departments(request: Request):
    await salary_user(request)
    rows = await db.salary_departments.find().sort("name", 1).to_list(500)
    counts = {}
    async for g in db.salary_workers.aggregate([{"$match": {"active": True}}, {"$group": {"_id": "$department", "n": {"$sum": 1}}}]):
        counts[g["_id"]] = g["n"]
    return [{**serialize_doc(r), "workers": counts.get(r["name"], 0)} for r in rows]


@router.post("/departments")
async def create_department(data: DeptIn, request: Request):
    user = await salary_user(request)
    name = data.name.strip()
    if await db.salary_departments.find_one({"name": {"$regex": f"^{re.escape(name)}$", "$options": "i"}}):
        raise HTTPException(400, f"A department called '{name}' already exists")
    res = await db.salary_departments.insert_one({"name": name, "created_at": _now()})
    note(request, entity_type="salary_department", entity_id=str(res.inserted_id), label=name, summary=f"Added salary department {name}")
    return {"id": str(res.inserted_id), "name": name, "workers": 0}


@router.delete("/departments/{dept_id}")
async def delete_department(dept_id: str, request: Request):
    await salary_user(request)
    if not ObjectId.is_valid(dept_id):
        raise HTTPException(404, "Department not found")
    dept = await db.salary_departments.find_one({"_id": ObjectId(dept_id)})
    if not dept:
        raise HTTPException(404, "Department not found")
    if await db.salary_workers.count_documents({"department": dept["name"]}):
        raise HTTPException(400, "Workers still belong to this department. Move them first.")
    note(request, entity_type="salary_department", entity_id=dept_id, label=dept["name"], summary=f"Deleted salary department {dept['name']}")
    await db.salary_departments.delete_one({"_id": dept["_id"]})
    return {"message": "Deleted"}


def present_worker(w):
    d = serialize_doc(dict(w))
    d["pay_type"] = d.get("pay_type") or "shift"
    d["hours_per_day"] = d.get("hours_per_day") if d["pay_type"] == "time" else None
    d["night_same_as_day"] = d.get("night_rate") is None
    d["night_rate"] = d["day_rate"] if d["night_same_as_day"] else d["night_rate"]
    return d


@router.get("/workers")
async def list_workers(request: Request, department: Optional[str] = None, q: Optional[str] = None, active: Optional[bool] = None):
    await salary_user(request)
    query = {}
    if department is not None:
        query["department"] = department
    if active is not None:
        query["active"] = active
    if q:
        rx = {"$regex": re.escape(q), "$options": "i"}
        query["$or"] = [{"name": rx}, {"emp_id": rx}, {"designation": rx}]
    rows = await db.salary_workers.find(query).sort("name", 1).to_list(5000)
    return [present_worker(r) for r in rows]


async def _dept_exists(name):
    if name and not await db.salary_departments.find_one({"name": name}):
        raise HTTPException(400, f"There is no department called '{name}'. Add it first.")


@router.post("/workers")
async def create_worker(data: WorkerIn, request: Request):
    await salary_user(request)
    await _dept_exists(data.department)
    if data.pay_type == "time" and not data.hours_per_day:
        raise HTTPException(400, "Time-based workers need their working hours per day")
    emp = data.emp_id.strip()
    if emp and await db.salary_workers.find_one({"emp_id": {"$regex": f"^{re.escape(emp)}$", "$options": "i"}}):
        raise HTTPException(400, f"Employee ID {emp} is already used")
    doc = {"name": data.name.strip(), "designation": data.designation.strip(), "emp_id": emp, "department": data.department,
           "day_rate": money(data.day_rate), "night_rate": money(data.night_rate) if data.night_rate is not None else None,
           "pay_type": data.pay_type, "hours_per_day": data.hours_per_day if data.pay_type == "time" else None, "active": True, "created_at": _now()}
    res = await db.salary_workers.insert_one(doc)
    doc["_id"] = res.inserted_id
    note(request, entity_type="salary_worker", entity_id=str(res.inserted_id), label=doc["name"], summary=f"Added salary worker {doc['name']}",
         after={"day_rate": doc["day_rate"], "night_rate": doc["night_rate"], "department": doc["department"], "pay_type": doc["pay_type"], "hours_per_day": doc["hours_per_day"]})
    return present_worker(doc)


@router.put("/workers/{worker_id}")
async def update_worker(worker_id: str, data: WorkerUpdate, request: Request):
    await salary_user(request)
    w = await _worker(worker_id)
    update = {k: v for k, v in data.model_dump().items() if v is not None}
    same = update.pop("night_same_as_day", None)
    if same:
        update["night_rate"] = None
        update.pop("night_rate_given", None)
    if not update:
        raise HTTPException(400, "No fields to update")
    if "pay_type" in update or "hours_per_day" in update:
        ptype = update.get("pay_type") or w.get("pay_type") or "shift"
        hpd = update.get("hours_per_day") if update.get("hours_per_day") is not None else w.get("hours_per_day")
        if ptype == "time" and not hpd:
            raise HTTPException(400, "Time-based workers need their working hours per day")
    if "name" in update:
        update["name"] = update["name"].strip()
    if "department" in update:
        await _dept_exists(update["department"])
    if "emp_id" in update:
        update["emp_id"] = update["emp_id"].strip()
        if update["emp_id"] and await db.salary_workers.find_one({"emp_id": {"$regex": f"^{re.escape(update['emp_id'])}$", "$options": "i"}, "_id": {"$ne": w["_id"]}}):
            raise HTTPException(400, f"Employee ID {update['emp_id']} is already used")
    for k in ("day_rate", "night_rate"):
        if update.get(k) is not None:
            update[k] = money(update[k])
    await db.salary_workers.update_one({"_id": w["_id"]}, {"$set": update})
    note(request, entity_type="salary_worker", entity_id=worker_id, label=w["name"], summary=f"Updated salary worker {w['name']}",
         before={k: w.get(k) for k in update}, after=update)
    return present_worker(await _worker(worker_id))


@router.delete("/workers/{worker_id}")
async def delete_worker(worker_id: str, request: Request):
    user = await salary_user(request)
    check_role(user, ["admin"])
    w = await _worker(worker_id)
    if await db.salary_entries.count_documents({"worker_id": worker_id}):
        raise HTTPException(400, "This worker has saved weeks, so they cannot be deleted. Switch them off instead.")
    note(request, entity_type="salary_worker", entity_id=worker_id, label=w["name"], summary=f"Deleted salary worker {w['name']}")
    await db.salary_workers.delete_one({"_id": w["_id"]})
    return {"message": "Deleted"}


async def _worker(worker_id):
    if not ObjectId.is_valid(worker_id):
        raise HTTPException(404, "Worker not found")
    w = await db.salary_workers.find_one({"_id": ObjectId(worker_id)})
    if not w:
        raise HTTPException(404, "Worker not found")
    return w


# ------------------------------------------------------------------ the week
class EntryIn(BaseModel):
    day_shifts: List[float] = Field(min_length=7, max_length=7)
    night_shifts: List[float] = Field(min_length=7, max_length=7)
    advance: float = Field(default=0, ge=0, le=10_000_000)
    notes: str = Field(default="", max_length=200)
    reprice: bool = False                      # use the worker's current rates instead of the ones saved with this week


def present_entry(e):
    d = serialize_doc(dict(e))
    return d


def totals(entries):
    return {"workers": len(entries), "gross": round(sum(e["gross"] for e in entries), 2), "advance": round(sum(e["advance"] for e in entries), 2),
            "net": round(sum(e["net"] for e in entries), 2), "day_shifts": round(sum(e.get("day_days", e["day_total"]) for e in entries), 2),
            "night_shifts": round(sum(e.get("night_days", e["night_total"]) for e in entries), 2)}


@router.get("/week")
async def get_week(request: Request, date: Optional[str] = None, start: Optional[str] = None):
    """Everything saved for one pay week. `date` may be any day; the Saturday starting its week is returned."""
    await salary_user(request)
    if start:
        s, e = week_bounds(start)
    else:
        d = parse_day(date) if date else datetime.now(timezone(timedelta(hours=5, minutes=30))).date()
        s = week_start_of(d).isoformat()
        s, e = week_bounds(s)
    rows = await db.salary_entries.find({"week_start": s}).sort("worker_name", 1).to_list(5000)
    entries = [present_entry(r) for r in rows]
    days = [(parse_day(s) + timedelta(days=i)).isoformat() for i in range(7)]
    return {"week_start": s, "week_end": e, "days": days, "day_names": DAYS, "entries": entries, "totals": totals(entries)}


@router.put("/week/{week_start}/worker/{worker_id}")
async def save_entry(week_start: str, worker_id: str, data: EntryIn, request: Request):
    user = await salary_user(request)
    s, e = week_bounds(week_start)
    w = await _worker(worker_id)
    existing = await db.salary_entries.find_one({"week_start": s, "worker_id": worker_id})
    if not existing and not w.get("active", True):
        raise HTTPException(400, f"{w['name']} is switched off. Switch them on in Workers before paying them.")
    if existing and not data.reprice:
        day_rate, night_rate = existing["day_rate"], existing["night_rate"]       # rates stay as paid
        ptype, hpd = existing.get("pay_type", "shift"), existing.get("hours_per_day")
    else:
        day_rate = w["day_rate"]
        night_rate = w.get("night_rate") if w.get("night_rate") is not None else w["day_rate"]
        ptype = w.get("pay_type") or "shift"
        hpd = w.get("hours_per_day") if ptype == "time" else None
    timed = ptype == "time"
    unit, top = ("hours", 24) if timed else ("shifts", 3)
    day = clean_shifts(data.day_shifts, "Day hours" if timed else "Day shift", QUARTER if timed else HALF, top)
    night = clean_shifts(data.night_shifts, "Night hours" if timed else "Night shift", QUARTER if timed else HALF, top)
    both = [i for i in range(7) if day[i] + night[i] > top]
    if both:
        raise HTTPException(400, f"{DAYS[both[0]]}: day and night {unit} together come to more than {top}. Check the numbers.")
    calc = compute(day, night, day_rate, night_rate, data.advance, hpd if timed else None)
    if calc["net"] < 0:
        raise HTTPException(400, f"The advance ({calc['advance']:,.2f}) is more than the wage ({calc['gross']:,.2f}). Reduce it; the rest can be taken next week.")
    doc = {"week_start": s, "week_end": e, "worker_id": worker_id, "worker_name": w["name"], "designation": w.get("designation", ""),
           "emp_id": w.get("emp_id", ""), "department": w.get("department", ""), "day_rate": day_rate, "night_rate": night_rate,
           "pay_type": ptype, "hours_per_day": hpd if timed else None, "day_shifts": day, "night_shifts": night, "notes": data.notes.strip(), **calc, "saved_by": user["id"], "saved_by_name": user.get("name", ""),
           "saved_at": _now()}
    before = None
    if existing:
        before = {"gross": existing["gross"], "advance": existing["advance"], "net": existing["net"], "day_total": existing["day_total"], "night_total": existing["night_total"]}
        await db.salary_entries.update_one({"_id": existing["_id"]}, {"$set": doc, "$push": {"history": {"at": _now(), "by": user.get("name", ""), "before": before}}})
        doc["_id"] = existing["_id"]
    else:
        doc["history"] = []
        doc["created_at"] = _now()
        res = await db.salary_entries.insert_one(doc)
        doc["_id"] = res.inserted_id
    note(request, entity_type="salary_entry", entity_id=str(doc["_id"]), label=f"{w['name']} · week of {s}",
         summary=f"{'Updated' if existing else 'Saved'} {w['name']}'s pay for the week of {s}: net {calc['net']:,.2f}",
         before=before, after={"gross": calc["gross"], "advance": calc["advance"], "net": calc["net"], "day_total": calc["day_total"], "night_total": calc["night_total"]})
    return present_entry(doc)


@router.delete("/week/{week_start}/worker/{worker_id}")
async def delete_entry(week_start: str, worker_id: str, request: Request):
    await salary_user(request)
    s, _ = week_bounds(week_start)
    e = await db.salary_entries.find_one({"week_start": s, "worker_id": worker_id})
    if not e:
        raise HTTPException(404, "Nothing saved for this worker in that week")
    note(request, entity_type="salary_entry", entity_id=str(e["_id"]), label=f"{e['worker_name']} · week of {s}",
         summary=f"Removed {e['worker_name']} from the week of {s}", before={"gross": e["gross"], "advance": e["advance"], "net": e["net"]})
    await db.salary_entries.delete_one({"_id": e["_id"]})
    return {"message": "Removed"}


class ClearIn(BaseModel):
    confirm: str


@router.post("/week/{week_start}/clear")
async def clear_week(week_start: str, data: ClearIn, request: Request):
    user = await salary_user(request)
    check_role(user, ["admin"])
    s, _ = week_bounds(week_start)
    if data.confirm != s:
        raise HTTPException(400, f"Type the week's start date ({s}) to confirm")
    rows = await db.salary_entries.find({"week_start": s}).to_list(5000)
    if not rows:
        raise HTTPException(404, "Nothing is saved for that week")
    # keep a copy: clearing a paid week is serious, and the audit log should be able to say what went
    await db.salary_cleared.insert_one({"week_start": s, "cleared_at": _now(), "cleared_by": user.get("name", ""), "entries": rows})
    await db.salary_entries.delete_many({"week_start": s})
    note(request, entity_type="salary_week", entity_id=s, label=f"Week of {s}", summary=f"Cleared the whole week of {s} ({len(rows)} workers)",
         before=totals(rows))
    return {"cleared": len(rows)}


@router.get("/weeks")
async def history(request: Request, page: int = Query(1, ge=1), page_size: int = Query(52, ge=1, le=200)):
    await salary_user(request)
    pipeline = [{"$group": {"_id": "$week_start", "workers": {"$sum": 1}, "gross": {"$sum": "$gross"}, "advance": {"$sum": "$advance"}, "net": {"$sum": "$net"}}},
                {"$sort": {"_id": -1}}]
    rows = [r async for r in db.salary_entries.aggregate(pipeline)]
    total = len(rows)
    rows = rows[(page - 1) * page_size: page * page_size]
    return {"items": [{"week_start": r["_id"], "week_end": (parse_day(r["_id"]) + timedelta(days=6)).isoformat(), "workers": r["workers"],
                       "gross": round(r["gross"], 2), "advance": round(r["advance"], 2), "net": round(r["net"], 2)} for r in rows],
            "total": total, "page": page, "pages": max(1, -(-total // page_size))}


# ------------------------------------------------------------------ Excel
def _safe(v):
    return ("'" + v) if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v


def build_xlsx(entries, weeks_summary=None, title="Salary"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    head = ["Week from", "Week to", "Emp ID", "Worker", "Designation", "Department", "Pay type", "Hours per day"] + [f"Day {d}" for d in DAYS] + ["Day shifts"] + [f"Night {d}" for d in DAYS] \
        + ["Night shifts", "Day rate", "Night rate", "Day wage", "Night wage", "Gross", "Advance", "Net pay", "Notes"]
    ws.append(head)
    for i in range(1, len(head) + 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="294F96", end_color="294F96", fill_type="solid")
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for e in entries:
        ws.append([e["week_start"], e["week_end"], _safe(e.get("emp_id", "")), _safe(e["worker_name"]), _safe(e.get("designation", "")), _safe(e.get("department", "")),
                   "By hours" if e.get("pay_type") == "time" else "By shift", e.get("hours_per_day") or ""]
                  + e["day_shifts"] + [e.get("day_days", e["day_total"])] + e["night_shifts"] + [e.get("night_days", e["night_total"])]
                  + [e["day_rate"], e["night_rate"], e["day_wage"], e["night_wage"], e["gross"], e["advance"], e["net"], _safe(e.get("notes", ""))])
    if entries:
        t = totals(entries)
        r = ws.max_row + 2
        ws.cell(row=r, column=head.index("Worker") + 1, value="TOTAL").font = Font(bold=True)
        sums = {"Day shifts": t["day_shifts"], "Night shifts": t["night_shifts"], "Day wage": round(sum(e["day_wage"] for e in entries), 2),
                "Night wage": round(sum(e["night_wage"] for e in entries), 2), "Gross": t["gross"], "Advance": t["advance"], "Net pay": t["net"]}
        for label, value in sums.items():
            ws.cell(row=r, column=head.index(label) + 1, value=value).font = Font(bold=True)
    ws.freeze_panes = "E2"
    for col, w in zip("ABCDEFGH", (12, 12, 9, 24, 16, 16, 10, 10)):
        ws.column_dimensions[col].width = w
    if weeks_summary:
        s2 = wb.create_sheet("Weeks")
        s2.append(["Week from", "Week to", "Workers", "Gross", "Advance", "Net pay"])
        for i in range(1, 7):
            s2.cell(row=1, column=i).font = Font(bold=True, color="FFFFFF")
            s2.cell(row=1, column=i).fill = PatternFill(start_color="294F96", end_color="294F96", fill_type="solid")
        for w in weeks_summary:
            s2.append([w["week_start"], w["week_end"], w["workers"], w["gross"], w["advance"], w["net"]])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


@router.get("/export")
async def export(request: Request, week: Optional[str] = None, all: bool = False):
    await salary_user(request)
    if all:
        rows = await db.salary_entries.find({}).sort([("week_start", 1), ("worker_name", 1)]).to_list(100000)
        summary = []
        async for r in db.salary_entries.aggregate([{"$group": {"_id": "$week_start", "workers": {"$sum": 1}, "gross": {"$sum": "$gross"}, "advance": {"$sum": "$advance"},
                                                                 "net": {"$sum": "$net"}}}, {"$sort": {"_id": 1}}]):
            summary.append({"week_start": r["_id"], "week_end": (parse_day(r["_id"]) + timedelta(days=6)).isoformat(), "workers": r["workers"],
                            "gross": round(r["gross"], 2), "advance": round(r["advance"], 2), "net": round(r["net"], 2)})
        name = "salary_all_weeks.xlsx"
        out = build_xlsx(rows, summary, "All weeks")
    else:
        if not week:
            raise HTTPException(400, "Give the week (its Saturday) or all=true")
        s, _ = week_bounds(week)
        rows = await db.salary_entries.find({"week_start": s}).sort("worker_name", 1).to_list(5000)
        if not rows:
            raise HTTPException(404, "Nothing is saved for that week")
        name = f"salary_{s}.xlsx"
        out = build_xlsx(rows, None, f"Week {s}")
    note(request, entity_type="salary_export", entity_id=week or "all", label="Salary export", summary=f"Exported salary to Excel ({name})")
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={name}"})


# ------------------------------------------------------------------ payslips
def _fmt_day(s):
    d = parse_day(s)
    return d.strftime("%a %d/%m/%Y")


def payslip_flowables(e, st):
    from reportlab.platypus import Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors as rl
    el = [Paragraph(f"<b>{COMPANY['name']}</b>", st["Title"]), Paragraph(COMPANY["address"], st["Normal"]), Spacer(1, 8),
          Paragraph(f"PAYSLIP · {_fmt_day(e['week_start'])} to {_fmt_day(e['week_end'])}", st["Heading2"]),
          Paragraph(f"<b>{e['worker_name']}</b>" + (f" · {e['designation']}" if e.get("designation") else "") + (f" · ID {e['emp_id']}" if e.get("emp_id") else "")
                    + (f" · {e['department']}" if e.get("department") else ""), st["Normal"]), Spacer(1, 8)]
    timed = e.get("pay_type") == "time"
    word = "hours" if timed else "shifts"
    grid = [["", *DAYS, "Total"], [f"Day {word}", *[f"{x:g}" for x in e["day_shifts"]], f"{e['day_total']:g}"],
            [f"Night {word}", *[f"{x:g}" for x in e["night_shifts"]], f"{e['night_total']:g}"]]
    t = Table(grid, colWidths=[70] + [38] * 7 + [48])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), rl.HexColor("#294f96")), ("TEXTCOLOR", (0, 0), (-1, 0), rl.white), ("GRID", (0, 0), (-1, -1), 0.4, rl.grey),
                           ("ALIGN", (1, 0), (-1, -1), "CENTER"), ("FONTSIZE", (0, 0), (-1, -1), 9)]))
    pay = [["", "Hours" if timed else "Shifts", "Rate", "Amount"],
           [f"Day {word}", f"{e['day_total']:g}", f"{e['day_rate']:,.2f}", f"{e['day_wage']:,.2f}"],
           [f"Night {word}", f"{e['night_total']:g}", f"{e['night_rate']:,.2f}", f"{e['night_wage']:,.2f}"],
           ["GROSS WAGE", "", "", f"{e['gross']:,.2f}"], ["Less: advance", "", "", f"-{e['advance']:,.2f}"], ["NET PAY", "", "", f"{e['net']:,.2f}"]]
    t2 = Table(pay, colWidths=[200, 70, 80, 100])
    t2.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), rl.HexColor("#294f96")), ("TEXTCOLOR", (0, 0), (-1, 0), rl.white), ("GRID", (0, 0), (-1, -1), 0.4, rl.grey),
                            ("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("FONTNAME", (0, 3), (-1, 3), "Helvetica-Bold"), ("FONTNAME", (0, 5), (-1, 5), "Helvetica-Bold"),
                            ("BACKGROUND", (0, 5), (-1, 5), rl.HexColor("#e8f0ea"))]))
    el += [t, Spacer(1, 12), t2]
    if timed:
        el += [Spacer(1, 4), Paragraph(f"Time-based: the rate is the salary for a day of {e['hours_per_day']:g} hours. Pay = hours worked / {e['hours_per_day']:g} x day salary.", st["Normal"])]
    if e.get("notes"):
        el += [Spacer(1, 6), Paragraph(f"Note: {e['notes'].replace('&', '&amp;').replace('<', '&lt;')}", st["Normal"])]
    el += [Spacer(1, 36), Paragraph("Signature: ___________________", st["Normal"])]
    return el


def build_pdf(entries, title):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, title=title)
    st = getSampleStyleSheet()
    el = []
    for i, e in enumerate(entries):
        if i:
            el.append(PageBreak())
        el += payslip_flowables(e, st)
    doc.build(el)
    out.seek(0)
    return out


@router.get("/payslip/{week_start}/{worker_id}")
async def payslip(week_start: str, worker_id: str, request: Request):
    await salary_user(request)
    s, _ = week_bounds(week_start)
    e = await db.salary_entries.find_one({"week_start": s, "worker_id": worker_id})
    if not e:
        raise HTTPException(404, "Nothing saved for this worker in that week")
    return StreamingResponse(build_pdf([e], f"Payslip {e['worker_name']}"), media_type="application/pdf",
                             headers={"Content-Disposition": f"inline; filename=payslip_{s}_{re.sub(r'[^A-Za-z0-9]+', '_', e['worker_name'])}.pdf"})


@router.get("/payslips/{week_start}")
async def payslips(week_start: str, request: Request):
    await salary_user(request)
    s, _ = week_bounds(week_start)
    rows = await db.salary_entries.find({"week_start": s}).sort("worker_name", 1).to_list(5000)
    if not rows:
        raise HTTPException(404, "Nothing is saved for that week")
    return StreamingResponse(build_pdf(rows, f"Payslips {s}"), media_type="application/pdf",
                             headers={"Content-Disposition": f"inline; filename=payslips_{s}.pdf"})
