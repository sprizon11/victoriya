"""
End-to-end check of the Victoria Unit Flow API.

Exercises every endpoint against a running backend and asserts the behaviour of
each bug fixed from DATA-ACCURACY-BUGS.md. Creates its own throwaway lot and
week so it does not disturb seeded demo data, and cleans up after itself.

    backend\\.venv\\Scripts\\python.exe backend\\test_api.py
"""

import io
import os
import sys
from datetime import datetime, timedelta
from pathlib import Path

import requests
from dotenv import load_dotenv

ROOT = Path(__file__).parent
load_dotenv(ROOT / ".env")

BASE = os.environ.get("TEST_BASE_URL", "http://127.0.0.1:8000")
EMAIL = os.environ.get("ADMIN_EMAIL", "admin@example.com").strip('"')
PASSWORD = os.environ.get("ADMIN_PASSWORD", "admin123").strip('"')

s = requests.Session()
passed, failed = [], []
created = {"lot": None, "week": None, "operators": [], "operation": None,
           "chart": None, "color": None, "advance": None}


def check(name, condition, detail=""):
    if condition:
        passed.append(name)
        print(f"  PASS  {name}")
    else:
        failed.append((name, detail))
        print(f"  FAIL  {name}   {detail}")


def api(method, path, **kw):
    return s.request(method, f"{BASE}/api{path}", timeout=30, **kw)


def section(title):
    print(f"\n--- {title} ---")


# ==========================================================================
section("auth")
r = api("post", "/auth/login", json={"email": EMAIL, "password": "wrong-password"})
check("bad password rejected", r.status_code == 401, f"got {r.status_code}")

r = api("post", "/auth/login", json={"email": EMAIL, "password": PASSWORD})
check("admin login", r.status_code == 200, f"got {r.status_code} {r.text[:120]}")
if r.status_code != 200:
    print("\nCannot continue without a login. Is the backend running?")
    sys.exit(1)

r = api("get", "/auth/me")
check("auth/me returns admin", r.status_code == 200 and r.json().get("role") == "admin")
check("password hash never returned", "password_hash" not in r.json())

bare = requests.get(f"{BASE}/api/lots", timeout=20)
check("unauthenticated request rejected", bare.status_code == 401, f"got {bare.status_code}")

# ==========================================================================
section("masters")
suffix = datetime.now().strftime("%H%M%S%f")[:10]

r = api("post", "/operations", json={"name": f"ZTest Op {suffix}"})
check("create operation", r.status_code == 200, r.text[:120])
created["operation"] = r.json().get("id")

r = api("post", "/operations", json={"name": f"ZTest Op {suffix}"})
check("duplicate operation rejected", r.status_code == 400)

r = api("post", "/colors", json={"name": f"ZTestColor{suffix}"})
check("create colour", r.status_code == 200)
created["color"] = r.json().get("id")
COLOR = r.json().get("name")

r = api("post", "/size-charts", json={"name": f"ZTest Chart {suffix}", "sizes": ["S", "M", "L"]})
check("create size chart", r.status_code == 200)
created["chart"] = r.json().get("id")

for i in range(2):
    r = api("post", "/operators", json={"name": f"ZTest Worker {suffix}-{i}"})
    check(f"create operator {i+1}", r.status_code == 200)
    created["operators"].append(r.json().get("id"))
check("operator codes auto-assigned",
      all(o for o in created["operators"]))

# ==========================================================================
section("weeks")
monday = datetime(2026, 3, 2)
r = api("post", "/weeks", json={"week_number": 901, "year": 2026,
                                "start_date": monday.strftime("%Y-%m-%d"),
                                "end_date": (monday + timedelta(days=6)).strftime("%Y-%m-%d")})
check("create week", r.status_code == 200, r.text[:120])
created["week"] = r.json().get("id")

r = api("post", "/weeks", json={"week_number": 901, "year": 2026,
                                "start_date": "2026-03-02", "end_date": "2026-03-08"})
check("duplicate week rejected", r.status_code == 400)

# ==========================================================================
section("lots")
RATE_A, RATE_B = 2.0, 3.0
lot_body = {
    "style_name": f"ZTEST-{suffix}",
    "start_date": "2026-03-02", "target_end_date": "2026-03-31",
    "size_chart_id": created["chart"],
    "color_names": [COLOR],
    "color_size_quantities": {COLOR: {"S": 100, "M": 100, "L": 100}},
    "operations": [
        {"operation_id": created["operation"], "rate": RATE_A,
         "operator_ids": created["operators"]},
    ],
}
r = api("post", "/lots", json=lot_body)
check("create lot", r.status_code == 200, r.text[:200])
lot = r.json()
created["lot"] = lot.get("id")
check("lot planned qty = 300", lot.get("total_planned_qty") == 300,
      f"got {lot.get('total_planned_qty')}")
check("new lot progress = 0", lot.get("progress") == 0)

# ==========================================================================
section("BUG 1 - daily entries must all save")
OP1, OP2 = created["operators"]
day1, day2 = "2026-03-02", "2026-03-03"

r = api("post", "/production/batch", json={
    "week_id": created["week"], "lot_id": created["lot"],
    "operation_id": created["operation"], "color": COLOR, "date": day1,
    "entries": [{"operator_id": OP1, "sizes": {"S": 10}}]})
check("day 1 entry saves", r.status_code == 200 and r.json()["saved"] == 1, r.text[:160])

r = api("post", "/production/batch", json={
    "week_id": created["week"], "lot_id": created["lot"],
    "operation_id": created["operation"], "color": COLOR, "date": day2,
    "entries": [{"operator_id": OP1, "sizes": {"S": 15}}]})
body = r.json() if r.status_code == 200 else {}
check("day 2 entry ALSO saves (was silently dropped)",
      r.status_code == 200 and body.get("saved") == 1 and not body.get("duplicates_skipped"),
      f"saved={body.get('saved')} skipped={body.get('duplicates_skipped')}")

r = api("post", "/production/batch", json={
    "week_id": created["week"], "lot_id": created["lot"],
    "operation_id": created["operation"], "color": COLOR, "date": day2,
    "entries": [{"operator_id": OP1, "sizes": {"S": 5}}]})
body = r.json() if r.status_code == 200 else {}
check("same operator same DAY is still a duplicate",
      r.status_code == 200 and body.get("saved") == 0 and body.get("duplicates_skipped"),
      f"saved={body.get('saved')}")

r = api("get", f"/production?lot_id={created['lot']}")
entries = r.json()
check("both days present, 25 pieces total",
      sum(e["total_quantity"] for e in entries) == 25,
      f"got {sum(e['total_quantity'] for e in entries)}")
check("earnings = 25 x rate",
      abs(sum(e["total_earnings"] for e in entries) - 25 * RATE_A) < 0.01)

# ==========================================================================
section("over-production guard")
r = api("post", "/production/batch", json={
    "week_id": created["week"], "lot_id": created["lot"],
    "operation_id": created["operation"], "color": COLOR, "date": "2026-03-04",
    "entries": [{"operator_id": OP2, "sizes": {"S": 999}}]})
check("over-plan batch rejected", r.status_code == 400, f"got {r.status_code}")
check("error names the size and totals", "S:" in r.text and "exceeds planned" in r.text)

# ==========================================================================
section("BUG 5 - edit must respect the plan")
entry_id = entries[0]["id"]
r = api("put", f"/production/{entry_id}", json={"size_quantities": {"S": 99999}})
check("over-plan EDIT rejected (guard was bypassable)", r.status_code == 400,
      f"got {r.status_code}")

r = api("put", f"/production/{entry_id}", json={"size_quantities": {"S": 12}})
check("valid edit accepted", r.status_code == 200, r.text[:160])
r = api("get", f"/production?lot_id={created['lot']}")
edited = next(e for e in r.json() if e["id"] == entry_id)
check("edit recomputes earnings", abs(edited["total_earnings"] - 12 * RATE_A) < 0.01,
      f"got {edited['total_earnings']}")
check("edit writes audit history", len(edited.get("edit_history", [])) >= 1)

r = api("put", f"/production/{entry_id}", json={"size_quantities": {"S": -5}})
check("negative quantity rejected", r.status_code == 400)

# ==========================================================================
section("BUG 2 - lot progress across operations")
# add a second operation, produce on it too, and confirm progress is not summed
lot_ops = lot["operations"] + []
r = api("get", "/operations")
other_op = next((o for o in r.json() if o["id"] != created["operation"]), None)
if other_op:
    r = api("put", f"/lots/{created['lot']}", json={"operations": [
        {"operation_id": created["operation"], "rate": RATE_A, "operator_ids": created["operators"]},
        {"operation_id": other_op["id"], "rate": RATE_B, "operator_ids": created["operators"]},
    ]})
    check("add second operation to lot", r.status_code == 200, r.text[:160])

    api("post", "/production/batch", json={
        "week_id": created["week"], "lot_id": created["lot"],
        "operation_id": other_op["id"], "color": COLOR, "date": day1,
        "entries": [{"operator_id": OP1, "sizes": {"S": 20}}]})

    r = api("get", f"/lots/{created['lot']}")
    L = r.json()
    # op A has 27 pieces (12 + 15), op B has 20 -> produced is the minimum, 20
    check("produced = min across operations, not the sum",
          L["total_produced_qty"] == 20,
          f"got {L['total_produced_qty']} (sum would be 47)")
    check("progress never exceeds 100%", L["progress"] <= 100, f"got {L['progress']}")
    check("balance is not negative", L["balance_qty"] >= 0, f"got {L['balance_qty']}")
    check("per-operation breakdown returned", len(L.get("operation_progress", [])) == 2)

# ==========================================================================
section("BUG 4 + 3 - payroll, advances and exports agree")
r = api("post", "/advances", json={"operator_id": OP1, "date": "2026-03-02",
                                   "amount": 100000, "remarks": "ZTest large advance"})
check("create advance", r.status_code == 200, r.text[:160])
created["advance"] = r.json().get("id")

r = api("post", "/advances", json={"operator_id": OP1, "date": "2026-03-02",
                                   "amount": -50, "remarks": "negative"})
check("negative advance rejected", r.status_code == 400)

r = api("get", "/payroll/weekly?week=901&year=2026")
check("payroll loads", r.status_code == 200, r.text[:160])
pay = r.json()
op1 = next((o for o in pay["operators"] if o["operator_id"] == OP1), None)
check("operator appears in payroll", op1 is not None)
if op1:
    check("advance deduction capped at earnings",
          op1["advance_deducted"] <= op1["total_earnings"] + 0.001,
          f"deducted {op1['advance_deducted']} vs earned {op1['total_earnings']}")
    check("net salary never negative", op1["net_salary"] >= 0,
          f"got {op1['net_salary']}")
    check("unrecovered advance carried forward", op1["advance_carried"] > 0)
    check("outstanding still visible", op1["advance_outstanding"] == 100000)
check("week net total not negative", pay["net_total"] >= 0, f"got {pay['net_total']}")

r = api("get", "/reports/export/excel?report_type=salary&week=901&year=2026")
check("salary Excel exports", r.status_code == 200, f"got {r.status_code}")
if r.status_code == 200:
    try:
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(r.content)).active
        text = " ".join(str(c) for row in ws.iter_rows(values_only=True) for c in row if c)
        check("salary sheet shows the advance deduction", "advance" in text.lower())
        check("salary sheet shows NET PAYABLE", "net payable" in text.lower())
    except Exception as e:
        check("salary sheet readable", False, str(e))

r = api("get", f"/reports/export/pdf?report_type=voucher&operator_id={OP1}&week=901&year=2026")
check("voucher PDF generates", r.status_code == 200 and r.content[:4] == b"%PDF",
      f"got {r.status_code}")
pdf_text = r.content
check("voucher PDF is non-trivial", len(pdf_text) > 1200, f"{len(pdf_text)} bytes")

# ==========================================================================
section("BUG 9 - voucher numbers are stable across reprints")
r1 = api("get", f"/reports/export/pdf?report_type=voucher&operator_id={OP1}&week=901&year=2026")
r2 = api("get", f"/reports/export/pdf?report_type=voucher&operator_id={OP1}&week=901&year=2026")
check("reprint returns same size (same number)", len(r1.content) == len(r2.content),
      f"{len(r1.content)} vs {len(r2.content)}")

# ==========================================================================
section("BUG 10 - lot Excel export")
r = api("get", f"/reports/export/excel?report_type=lot&lot_id={created['lot']}")
check("lot Excel export no longer 500s", r.status_code == 200, f"got {r.status_code}")
check("lot Excel has content", len(r.content) > 3000, f"{len(r.content)} bytes")

r = api("get", f"/reports/export/excel?report_type=operation_cost&lot_id={created['lot']}")
check("operation cost export", r.status_code == 200)

r = api("get", "/reports/export/pdf?report_type=salary&week=901&year=2026")
check("salary PDF export", r.status_code == 200 and r.content[:4] == b"%PDF")

# ==========================================================================
section("BUG 8 - advance adjustment honours the chosen voucher")
r = api("post", "/advances", json={"operator_id": OP2, "date": "2026-01-05",
                                   "amount": 500, "remarks": "ZTest older"})
older = r.json().get("id")
r = api("post", "/advances", json={"operator_id": OP2, "date": "2026-02-05",
                                   "amount": 700, "remarks": "ZTest newer"})
newer = r.json().get("id")

r = api("post", "/payroll/adjust-advance", json={
    "advance_id": newer, "operator_id": OP2, "amount": 200,
    "week_number": 901, "year": 2026})
check("adjust selected advance", r.status_code == 200, r.text[:160])

r = api("get", f"/advances?operator_id={OP2}")
advs = {a["id"]: a for a in r.json()}
check("the SELECTED voucher was reduced", advs[newer]["balance"] == 500,
      f"newer balance {advs[newer]['balance']}")
check("the other voucher was left alone", advs[older]["balance"] == 500,
      f"older balance {advs[older]['balance']}")

r = api("post", "/payroll/adjust-advance", json={"operator_id": OP2, "amount": 0,
                                                 "week_number": 901, "year": 2026})
check("zero adjustment rejected", r.status_code == 400)
r = api("post", "/payroll/adjust-advance", json={"operator_id": OP2, "amount": 999999,
                                                 "week_number": 901, "year": 2026})
check("over-balance adjustment rejected", r.status_code == 400)

# ==========================================================================
section("BUG 6 - edit requests actually apply")
r = api("post", "/edit-requests", json={
    "entry_type": "production", "entry_id": entry_id,
    "reason": "ZTest miscount", "proposed_changes": {"size_quantities": {"S": 7}}})
check("create edit request", r.status_code == 200, r.text[:160])
req_id = r.json().get("id")

r = api("put", f"/edit-requests/{req_id}", json={"action": "approve"})
check("approve edit request", r.status_code == 200, r.text[:200])

r = api("get", f"/production?lot_id={created['lot']}")
after = next(e for e in r.json() if e["id"] == entry_id)
check("approval CHANGED the entry (was a no-op)", after["total_quantity"] == 7,
      f"got {after['total_quantity']}")
check("approval recomputed earnings", abs(after["total_earnings"] - 7 * RATE_A) < 0.01)

r = api("put", f"/edit-requests/{req_id}", json={"action": "approve"})
check("double-approval rejected", r.status_code == 400)

r = api("post", "/edit-requests", json={
    "entry_type": "production", "entry_id": entry_id, "reason": "no changes attached"})
empty_req = r.json().get("id")
r = api("put", f"/edit-requests/{empty_req}", json={"action": "approve"})
check("approving an empty request is refused, not silently ignored", r.status_code == 400)

# ==========================================================================
section("BUG 12 - lot deletion cannot rewrite paid history")
r = api("delete", f"/lots/{created['lot']}")
check("delete refused while entries exist", r.status_code == 400, f"got {r.status_code}")
check("refusal explains why", "production entries" in r.text.lower())

# ==========================================================================
section("reports")
for path in [f"/reports/lot-matrix/{created['lot']}",
             "/reports/weekly-matrix?week=901&year=2026",
             "/reports/weekly-production?week=901&year=2026",
             f"/reports/lot-color/{created['lot']}",
             f"/reports/lot-operation/{created['lot']}",
             f"/reports/operator-detail/{OP1}?week=901&year=2026",
             f"/reports/lot/{created['lot']}",
             "/reports/advance-report?week=901&year=2026"]:
    r = api("get", path)
    check(f"GET {path.split('?')[0]}", r.status_code == 200, f"got {r.status_code}")

# ==========================================================================
section("float precision")
r = api("get", "/dashboard/summary")
check("dashboard loads", r.status_code == 200)
dash = r.json()


def clean(v):
    return round(float(v), 2) == float(v)


check("dashboard week earnings rounded", clean(dash["week_earnings"]),
      f"got {dash['week_earnings']}")
check("top operator earnings rounded",
      all(clean(o["total_earnings"]) for o in dash["top_operators"]),
      str([o["total_earnings"] for o in dash["top_operators"]][:4]))
check("operation stats rounded",
      all(clean(o["total_earnings"]) for o in dash["operation_stats"]))

r = api("get", "/reports/weekly-production?week=30&year=2026")
if r.status_code == 200:
    wp = r.json()
    check("weekly report lot earnings rounded",
          all(clean(l["total_earnings"]) for l in wp["lots"]),
          str([l["total_earnings"] for l in wp["lots"]][:4]))

# ==========================================================================
section("misc endpoints")
for path in ["/company", "/health", "/operators", "/operations", "/colors",
             "/size-charts", "/weeks", "/lots", "/advances", "/users",
             "/settings", "/notifications", "/notifications/unread-count",
             "/vouchers?week=30&year=2026", "/edit-requests"]:
    r = api("get", path)
    check(f"GET {path.split('?')[0]}", r.status_code == 200, f"got {r.status_code}")

# ==========================================================================
section("cleanup")
api("delete", f"/production/{entry_id}")
for e in api("get", f"/production?lot_id={created['lot']}").json():
    api("delete", f"/production/{e['id']}")
r = api("delete", f"/lots/{created['lot']}")
check("lot deletes once entries are gone", r.status_code == 200, r.text[:160])
for aid in [created["advance"], older, newer]:
    if aid:
        api("delete", f"/advances/{aid}")
for oid in created["operators"]:
    api("delete", f"/operators/{oid}")
api("delete", f"/operations/{created['operation']}")
api("delete", f"/colors/{created['color']}")
api("delete", f"/size-charts/{created['chart']}")
api("delete", f"/weeks/{created['week']}")
print("  cleaned up test fixtures")

# ==========================================================================
print("\n" + "=" * 60)
print(f"  {len(passed)} passed, {len(failed)} failed")
if failed:
    print("\n  FAILURES:")
    for name, detail in failed:
        print(f"    - {name}  {detail}")
print("=" * 60)
sys.exit(1 if failed else 0)
